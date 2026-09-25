"""Recompute annotation agreement directly from the two released workbooks."""

import argparse
from collections import Counter, defaultdict
import hashlib
import json
from pathlib import Path
import statistics

import openpyxl


POOLS = ("Natural500", "Final500", "StressFresh100")
CODERS = ("coder_01", "coder_02")
STEP_LABELS = {"有进展": "progress", "停滞": "stagnant", "不确定": "uncertain"}
SPAN_LABELS = {"有效迭代": "productive_iteration", "无效循环": "unproductive_cycle"}


def require(condition, message):
    if not condition:
        raise ValueError(message)


def agreement(left, right):
    require(len(left) == len(right) and len(left) > 0, "Empty or unequal paired samples")
    n = len(left)
    labels = sorted(set(left) | set(right))
    a, b = Counter(left), Counter(right)
    matched = sum(x == y for x, y in zip(left, right))
    observed = matched / n
    expected = sum(a[label] * b[label] for label in labels) / (n * n)
    return {
        "paired_records": n,
        "agreeing_records": matched,
        "raw_agreement": observed,
        "cohens_kappa": (observed - expected) / (1 - expected) if expected != 1 else None,
        "coder_01_counts": {label: a[label] for label in labels},
        "coder_02_counts": {label: b[label] for label in labels},
        "confusion_matrix": {
            x: {y: sum(u == x and v == y for u, v in zip(left, right)) for y in labels}
            for x in labels
        },
    }


def read_workbook(path):
    result = {}
    with path.open("rb") as handle:
        workbook = openpyxl.load_workbook(handle, read_only=True, data_only=False)
        for pool in POOLS:
            steps = {}
            indices = defaultdict(list)
            for row_number, row in enumerate(workbook[pool + "步骤"].iter_rows(min_row=2, values_only=True), 2):
                require(len(row) == 5, f"{pool}: unexpected step columns")
                dataset, trajectory, index, label, confidence = row
                require(dataset == pool, f"{pool}: dataset mismatch at step row {row_number}")
                require(type(index) is int and index >= 0, f"{pool}: invalid step index")
                require(label in STEP_LABELS, f"{pool}: invalid step label")
                require(confidence is None or confidence in (1, 2, 3, 4, 5), f"{pool}: invalid confidence")
                key = (trajectory, index)
                require(key not in steps, f"{pool}: duplicate step key {key}")
                steps[key] = STEP_LABELS[label]
                indices[trajectory].append(index)
            for trajectory, values in indices.items():
                require(sorted(values) == list(range(max(values) + 1)), f"{pool}: noncontiguous steps in {trajectory}")
            spans = []
            for row_number, row in enumerate(workbook[pool + "片段"].iter_rows(min_row=2, values_only=True), 2):
                require(len(row) == 7, f"{pool}: unexpected span columns")
                dataset, trajectory, sequence, start, end, label, confidence = row
                require(dataset == pool and trajectory in indices, f"{pool}: invalid span trajectory")
                require(type(start) is int and type(end) is int, f"{pool}: noninteger boundaries")
                require(0 <= start <= end <= max(indices[trajectory]), f"{pool}: span outside trajectory")
                require(label in SPAN_LABELS, f"{pool}: invalid span label")
                require(confidence is None or confidence in (1, 2, 3, 4, 5), f"{pool}: invalid confidence")
                spans.append({"trajectory": trajectory, "start": start, "end": end,
                              "label": SPAN_LABELS[label], "row": row_number})
            result[pool] = {"steps": steps, "spans": spans, "trajectories": set(indices)}
        workbook.close()
    return result


def match_spans(left, right):
    """Match without consulting labels; use original row order to break IoU ties."""
    a, b = defaultdict(list), defaultdict(list)
    for span in left:
        a[span["trajectory"]].append(span)
    for span in right:
        b[span["trajectory"]].append(span)
    pairs = []
    for trajectory in sorted(set(a) | set(b)):
        candidates = []
        for i, x in enumerate(a[trajectory]):
            for j, y in enumerate(b[trajectory]):
                intersection = max(0, min(x["end"], y["end"]) - max(x["start"], y["start"]) + 1)
                union = x["end"] - x["start"] + 1 + y["end"] - y["start"] + 1 - intersection
                iou = intersection / union
                if iou >= 0.5:
                    candidates.append((-iou, i, j))
        used_a, used_b = set(), set()
        for _, i, j in sorted(candidates):
            if i in used_a or j in used_b:
                continue
            used_a.add(i)
            used_b.add(j)
            pairs.append((a[trajectory][i], b[trajectory][j]))
    return pairs


def span_statistics(pairs, count_a, count_b):
    result = agreement([a["label"] for a, _ in pairs], [b["label"] for _, b in pairs])
    identical = sum(a["start"] == b["start"] and a["end"] == b["end"] for a, b in pairs)
    result.update({
        "coder_01_spans": count_a,
        "coder_02_spans": count_b,
        "coder_01_matched_fraction": len(pairs) / count_a,
        "coder_02_matched_fraction": len(pairs) / count_b,
        "coder_01_unmatched_spans": count_a - len(pairs),
        "coder_02_unmatched_spans": count_b - len(pairs),
        "identical_endpoint_pairs": identical,
        "identical_endpoint_fraction": identical / len(pairs),
        "median_start_absolute_deviation": statistics.median(abs(a["start"] - b["start"]) for a, b in pairs),
        "median_end_absolute_deviation": statistics.median(abs(a["end"] - b["end"]) for a, b in pairs),
    })
    return result


def reproduce(root):
    files = {coder: root / coder / "标注记录.xlsx" for coder in CODERS}
    books = {coder: read_workbook(path) for coder, path in files.items()}
    result = {"workbook_sha256": {str(path.relative_to(root)).replace("\\", "/"): hashlib.sha256(path.read_bytes()).hexdigest()
                                   for path in files.values()}, "pools": {}}
    all_a, all_b, all_pairs = [], [], []
    count_a = count_b = 0
    for pool in POOLS:
        a, b = (books[coder][pool] for coder in CODERS)
        require(set(a["steps"]) == set(b["steps"]), f"{pool}: unequal step keys")
        keys = sorted(a["steps"])
        aa, bb = [a["steps"][key] for key in keys], [b["steps"][key] for key in keys]
        pairs = match_spans(a["spans"], b["spans"])
        result["pools"][pool] = {
            "trajectories": len(a["trajectories"]),
            "step_agreement": agreement(aa, bb),
            "span_agreement": span_statistics(pairs, len(a["spans"]), len(b["spans"])),
            "trajectories_without_spans": {
                coder: sorted(books[coder][pool]["trajectories"] - {s["trajectory"] for s in books[coder][pool]["spans"]})
                for coder in CODERS
            },
        }
        all_a.extend(aa)
        all_b.extend(bb)
        all_pairs.extend(pairs)
        count_a += len(a["spans"])
        count_b += len(b["spans"])
    result["overall_step_agreement"] = agreement(all_a, all_b)
    result["step_one_vs_rest_kappa"] = {
        label: agreement([x == label for x in all_a], [x == label for x in all_b])["cohens_kappa"]
        for label in sorted(set(all_a) | set(all_b))
    }
    result["uncertain_steps"] = {"coder_01": all_a.count("uncertain"), "coder_02": all_b.count("uncertain"),
                                 "both": sum(x == y == "uncertain" for x, y in zip(all_a, all_b)),
                                 "either": sum(x == "uncertain" or y == "uncertain" for x, y in zip(all_a, all_b))}
    result["overall_span_agreement"] = span_statistics(all_pairs, count_a, count_b)
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true", help="Compare recomputed results with agreement_summary.json")
    parser.add_argument("--output", type=Path, help="Write recomputed JSON to a chosen path")
    args = parser.parse_args()
    root = Path(__file__).resolve().parent
    result = reproduce(root)
    if args.check:
        expected = json.loads((root / "agreement_summary.json").read_text(encoding="utf-8"))
        require(result == expected, "Recomputed results differ from agreement_summary.json")
        print("PASS: workbook hashes, counts, and all reported agreement statistics reproduced.")
    if args.output:
        args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
    if not args.check and not args.output:
        print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
